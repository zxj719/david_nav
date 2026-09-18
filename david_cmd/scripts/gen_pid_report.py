#!/usr/bin/env python3
# Copyright 2026 xingjian zhang
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""根据最近几组 PI 阶跃实验生成 Word 调参报告（含阶跃响应图 + 参数/结果对比 + 分析）."""
import glob
import importlib.util
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

from docx import Document  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location("rag", os.path.join(HERE, "run_a_group.py"))
rag = importlib.util.module_from_spec(spec)
spec.loader.exec_module(rag)
analyze_step, parse_kpki, _avg = rag.analyze_step, rag.parse_kpki, rag._avg

BASE = "/home/agx/a_group_results"
FIGDIR = os.path.join(BASE, "_pid_report_figs")
os.makedirs(FIGDIR, exist_ok=True)

# (目录后缀, 说明, axis, target, 单位)
RUNS = [
    ("214545", "Ki=0 扫 Kp", "angular", 0.3, "rad/s"),
    ("215353", "降增益 + 等比 Ki / 高 Ki", "angular", 0.2, "rad/s"),
    ("220120", "固定 Kp 扫 Ki", "angular", 0.3, "rad/s"),
    ("220650", "revert 直行验证", "linear", 0.12, "m/s"),
    ("220812", "revert 高速验证", "angular", 0.5, "rad/s"),
]


def dpath(suf):
    return os.path.join(BASE, "20260608_" + suf)


def labels_of(suf):
    fs = glob.glob(os.path.join(dpath(suf), "vle", "PI_*_pos.csv"))
    labs = [os.path.basename(f)[3:-8] for f in fs]

    def key(label):
        kp, ki = parse_kpki(label)
        try:
            return (float(kp), float(ki) if ki != "-" else -1)
        except ValueError:
            return (0, 0)
    return sorted(labs, key=key)


def trace(suf, label, axis):
    col = 2 if axis == "angular" else 1
    path = os.path.join(dpath(suf), "vle", "PI_%s_pos.csv" % label)
    t, v = [], []
    with open(path) as fh:
        for i, line in enumerate(fh):
            if i == 0:
                continue
            p = line.strip().split(",")
            if len(p) > col:
                t.append(float(p[0]))
                v.append(float(p[col]))
    if t:
        t0 = t[0]
        t = [x - t0 for x in t]
    return t, v


def metric_row(suf, label, axis):
    vle = os.path.join(dpath(suf), "vle")
    mp = analyze_step(os.path.join(vle, "PI_%s_pos.csv" % label), axis=axis)
    mn = analyze_step(os.path.join(vle, "PI_%s_neg.csv" % label), axis=axis)
    kp, ki = parse_kpki(label)
    g = lambda k: _avg(mp, mn, k)  # noqa: E731
    return [kp, ki, g("sse_pct"), g("overshoot_pct"), g("rise_s"),
            g("settle_s"), g("reversal"), g("osc_crossings")]


def fmt(x, nd=2):
    if x is None:
        return "n/a"
    try:
        return ("%." + str(nd) + "f") % float(x)
    except (TypeError, ValueError):
        return str(x)


# ---------------- 生成阶跃响应图 ----------------
def fig_overlay(suf, labels, axis, target, unit, title, fname):
    plt.figure(figsize=(6.2, 3.4))
    for lab in labels:
        t, v = trace(suf, lab, axis)
        plt.plot(t, v, linewidth=1.0, label=lab.replace("_", " "))
    plt.axhline(target, color="k", linestyle="--", linewidth=0.8,
                label="target %.2f" % target)
    plt.xlabel("time (s)")
    plt.ylabel("odom %s (%s)" % ("wz" if axis == "angular" else "vx", unit))
    plt.title(title)
    plt.legend(fontsize=7)
    plt.grid(alpha=0.3)
    plt.tight_layout()
    p = os.path.join(FIGDIR, fname)
    plt.savefig(p, dpi=130)
    plt.close()
    return p


figs = {}
figs["ki0"] = fig_overlay("214545", labels_of("214545"), "angular", 0.3, "rad/s",
                          "Ki=0 sweep @0.3 rad/s : all Kp limit-cycle", "f_ki0.png")
figs["low"] = fig_overlay("215353", ["Kp75_Ki75", "Kp300_Ki300"], "angular", 0.2,
                          "rad/s", "Low gain stalls vs mid gain stuck-then-lurch @0.2",
                          "f_low.png")
figs["ki"] = fig_overlay("220120", ["Kp300_Ki300", "Kp300_Ki800"], "angular", 0.3,
                         "rad/s", "Higher Ki breaks free sooner @0.3", "f_ki.png")
figs["rev"] = fig_overlay("220812", ["Kp300_Ki800"], "angular", 0.5, "rad/s",
                          "Revert Kp300_Ki800 @0.5 : startup lurch ~2.3x then wander",
                          "f_rev.png")

# ---------------- 组装 Word ----------------
doc = Document()
doc.styles["Normal"].font.size = Pt(10.5)


def h(text, lvl=1):
    doc.add_heading(text, level=lvl)


def p(text, bold=False):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold = bold
    return par


def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, hd in enumerate(headers):
        t.rows[0].cells[i].text = hd
    for r in rows:
        c = t.add_row().cells
        for i, val in enumerate(r):
            c[i].text = str(val)
    return t


title = doc.add_heading("重载差速底盘 增量式速度环 PI 调参报告", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph("对象：70 kg 上肢 + 差速底盘 + 被动万向轮（David）｜日期：2026-06-08")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

h("1. 概述与核心结论", 1)
p("本报告基于 2026-06-08 在 2×2 m 场地对底盘速度环做的 5 组阶跃响应实验，"
  "目的是为手动 flash 进 STM32 的增量式速度 PI 选定合适参数。核心结论：", )
bullet("该底盘低速段存在严重 stick-slip（静摩擦≫动摩擦）+ 被动万向轮破启阻力，"
       "速度环在任何测试增益下都无法获得干净的低速跟踪——这是物理特性，不是参数没调好。")
bullet("三个增益区间各有死法：低增益→底盘卡死不动；中增益（含原始 300/300）→憋停约 2 s 后暴力挣脱、永不整定；"
       "高增益（1200）→持续极限环振荡（±65%）。")
bullet("折中 revert 值：Kp=300 / Ki=800。它在正常速度下平均能跟、且没有 1200 那种持续极限环，"
       "严格优于当前的 1200，建议先 flash 它把车从极限环中救出。但它不是“解”，起步暴冲与低速顿挫依旧。")
bullet("治本方向：转每轮位置（增量）控制 + 按上一段旋转量 Δθ 自适应的破启前馈脉冲；"
       "并在导航层避免进入 stick-slip 低速区。")

h("2. 实验方法与环境", 1)
bullet("指令注入：直接发布 /cmd_vel（geometry_msgs/Twist），绕过 david_move/cmd_vel_separator，测电机+底盘裸特性。")
bullet("反馈：仅轮速计 /odom（chassis 级 linear.x / angular.z），无单轮编码器、未用 IMU。")
bullet("激励：阶跃（velocity_loop_eval --mode step），长 hold（角 4 s / 直行 3 s）确保观察到平台或振荡全貌；每组取 +/- 两方向均值。")
bullet("控制器：STM32 增量式 PI，Kp/Ki 由人工逐组 flash；已有抗饱和（增量式本身几乎不积分饱和，故“反冲”列主要是停车主动制动过冲，非 windup）。")
bullet("场地：2×2 m，故线速度只能测到低速档（高速/大半径未覆盖）。")

h("2.1 指标定义", 2)
table(["指标", "含义", "理想"],
      [["稳态误差%", "1 − 稳态实速/目标（负值=冲过目标）", "<5%"],
       ["超调%", "(峰值−稳态)/稳态；重载下多为脚轮挣脱尖峰", "<10–15%"],
       ["上升s", "10%→90% 上升时间", "小"],
       ["整定s", "进入 ±5% 带的时间；≈hold 即“未整定/持续振荡”", "<1.2 s"],
       ["反冲", "命令归零后反向速度（主动制动过冲）", "≈0"],
       ["振荡", "稳态区过零次数（极限环指标）", "0"]])

h("3. 参数与结果对比", 1)
for suf, desc, axis, tgt, unit in RUNS:
    labs = labels_of(suf)
    if not labs:
        continue
    p("实验 %s ｜ %s ｜ %s 阶跃，目标 %.2f %s" % (suf, desc, axis, tgt, unit), bold=True)
    rows = []
    for lab in labs:
        m = metric_row(suf, lab, axis)
        hold = 4.0 if axis == "angular" else 3.0
        unsettled = m[5] is not None and m[5] >= 0.8 * hold
        rows.append([lab, m[0], m[1], fmt(m[2], 1), fmt(m[3], 1), fmt(m[4]),
                     fmt(m[5]) + ("⚠" if unsettled else ""),
                     fmt(m[6], 3), fmt(m[7], 1)])
    table(["PI 组", "Kp", "Ki", "稳态误差%", "超调%", "上升s", "整定s", "反冲", "振荡"], rows)
    doc.add_paragraph()

h("4. 阶跃响应曲线", 1)
cap = [
    (figs["ki0"], "图1 Ki=0 时各 Kp（300/600/900/1200）@0.3 rad/s：全部持续极限环振荡，"
     "整个 hold 不收敛。证明纯比例（增量式 Ki=0）在该 plant 上不可用。"),
    (figs["low"], "图2 @0.2 rad/s：Kp75/Ki75（蓝）全程≈0——底盘被静摩擦完全卡死；"
     "Kp300/Ki300（橙，原始值）憋停约 2 s 后暴力挣脱冲高。说明“降增益”会卡死，此路不通。"),
    (figs["ki"], "图3 @0.3 rad/s：相同 Kp=300 下，Ki=800（橙）比 Ki=300（蓝）更早挣脱、"
     "更接近目标——高 Ki 帮助 punch 穿 stiction。"),
    (figs["rev"], "图4 revert 值 Kp300/Ki800 @0.5 rad/s：起步卡 ~0.5 s 后猛冲到约 1.17（目标 0.5，约 2.3 倍），"
     "随后在 0.3–0.7 间游荡。即正常速度也有强起步暴冲（=脚轮破启）+ 残余 stick-slip 游荡。"),
]
for img, c in cap:
    doc.add_picture(img, width=Inches(5.6))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp = doc.add_paragraph(c)
    cp.runs[0].font.size = Pt(9)
    cp.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

h("5. 分析与解释", 1)
h("5.1 为什么速度环“调不出来”——三区间夹击", 2)
p("阶跃曲线一致表明：底盘只有“卡死”和“挣脱后猛冲”两个状态，中间的“平稳低速”区在该车上物理上不存在：")
bullet("低增益（Kp≤150）：输出不足以克服静摩擦+脚轮破启阻力 → 底盘根本不动（图2 蓝）。")
bullet("中增益（原始 Kp300/Ki300）：憋停约 2 s，积分累积到挣脱后暴冲、过冲到 2–3 倍、永不整定（图2 橙、图4）。")
bullet("高增益（Kp1200）：挣脱后速度环反复“冲过头→反修→再粘住”，形成持续极限环（图1）。")
p("结论：这不是 PI 没调好，而是速度环无法驯服该 plant 的 stick-slip + 大破启。")

h("5.2 增量式 PI 下 Kp / Ki 的作用", 2)
bullet("增量式与位置式数学等价（Kp/Ki 角色相同）：Ki·e 提供把输出拉向目标的恢复/消稳差作用；Kp·Δe 提供阻尼。")
bullet("故 Ki=0 = 纯比例，缺乏恢复项，叠加大死时间 → 必然振荡（图1）。这也说明此前“先 Ki=0 扫 Kp”是套了位置式教条、不适合本 plant。")
bullet("相同 Kp 下提高 Ki 反而更温和（图3）：高 Ki 让底盘更快挣脱，不至于长时间憋停后暴冲。")

h("5.3 起步暴冲 = 被动万向轮破启", 2)
p("每组阶跃起步都先卡死一段再猛挣脱，且每次卡的时长接近一致——因为每次起步时脚轮都停在同一个不利（反向）朝向，"
  "需先甩正、期间横向刮擦产生大破启阻力。0.5 rad/s 时这记挣脱冲到约 1.17 rad/s（图4），随速度增大而增大。"
  "即“起步段由脚轮破启主导，挣脱后的持续游荡由速度环 stick-slip 主导”。")

h("6. 结论与建议", 1)
bullet("立即：将 STM32 速度 PI 从 1200 退回 Kp=300 / Ki=800（least-bad，正常速度可用、无持续极限环）。")
bullet("放弃用速度环追求低速平稳——与物理较劲；低速顿挫是该底盘固有。")
bullet("治本一：转每轮位置（增量）控制——位置环处理破启更优雅（误差累积→挣脱→追上），且不像速度环持续极限环；"
       "需配 50 Hz 流式位置 setpoint + 平滑位置斜坡。")
bullet("治本二：破启前馈脉冲，按上一段旋转量 Δθ 自适应（脚轮越歪 kick 越大），位置/速度模式均适用。")
bullet("导航层：曲率连续轨迹 + 航向闭环 + 避免进入 stick-slip 低速区（用 kick 穿透、精定位用短位置脉冲），"
       "这与现有 cmd_vel_separator 的设计方向一致。")
bullet("后续验证：在 2×2 内用 /odom 位姿做位置斜坡跟踪 + 落点精度实验，量化位置环相对速度环的改善。")

h("7. 速度死区验证（最终 PI 下）", 1)
p("在最终 revert 参数 Kp=300/Ki=800、直发 /cmd_vel、仅轮速计判定下，对低速档做破启扫描。"
  "因脚轮 70 kg 推不动，用机器人自驱预置脚轮；下表按“最坏（逆向）脚轮”取值。")
table(["量", "裸破启命令", "该档实速", "生产 lift (cmd.yaml)", "单位"],
      [["线速度", "≈0.04", "0.023（爬行）", "min_linear_x = 0.20", "m/s"],
       ["角速度", "≈0.15", "0.037（爬行）", "min_stationary_angular_z = 0.35", "rad/s"]])
p("交叉印证：PI=1200 时测得 0.035 / 0.12；本次 Kp300/Ki800（增益更低）为 0.04 / 0.15。"
  "→ 死区随增益降低而增大，证明它是“增益 × 脚轮朝向”的耦合量，并非定值。", )
h("7.1 三个层次的速度，别混", 2)
bullet("裸破启（≈0.04 / 0.15）：仅“抽动一下”，该档实速只有 0.023 / 0.037 且 stick-slip 抖动，不是干净运动。")
bullet("可用速度（未实测，介于二者之间）：能持续、干净运动的最低命令。应看稳态 ss，而非破启峰值。")
bullet("生产 lift（0.20 / 0.35）：保证任意脚轮态、过桥后可靠不抖起步——比裸破启大 2.3–5× 是合理裕度，非浪费。")
p("结论：生产值偏保守但合理；若要为低速平顺而下调，应专门测“最低可持续干净运动”的命令，不能用裸破启 0.04/0.15。", bold=True)
p("数据说明：本次死区取自来回 amp 扫描（该方式每段会让脚轮翻向，仅首段为预置态），故为近似最坏值；"
  "脚本已新增 dz 模式（每档重新 precond + 单方向）可获得干净的顺/逆向一对。")

h("8. 复盘：为什么没能一次分析到位", 1)
p("从最初“万向轮拖尾/循迹差”到最终“速度环低速不可控 + 死区验证”，经历了多轮迭代。原因分两类——我的失误，"
  "和问题本身的客观难度。如实记录，供后续少走弯路。")
h("8.1 我的失误", 2)
bullet("过早给方案而非先测 plant：先抛了“电流/力矩前馈”，但底盘只有速度接口、根本无法注入力——应在动笔前先确认控制接口边界。")
bullet("套用通用教条到特殊对象：按位置式 PID 直觉建议“先 Ki=0 扫 Kp”，但这是增量式 PI（Ki=0=纯比例），叠加大死时间必然振荡，白跑一轮。")
bullet("过信汇总指标、没先看波形：第一轮 hold 仅 1 s，速度没到稳态就采样，得出“增益>1”的伪值并一度误导；回看原始 CSV 波形才发现是“还在上升的斜坡”。")
bullet("指标语义未对齐物理：把 stick-slip 的“挣脱尖峰”当成控制“超调”；自动报告早期还在“全部不合格”时给出误导性“推荐 Kp1200”（缺少不合格门槛，后补）。")
bullet("实验设计考虑不周，被用户两次纠正：①预置脚轮假设能手推（70 kg 推不动）；②来回 amp 扫描会让脚轮每段翻向、污染死区测量。")
bullet("工具选择反复：A1 最初用强依赖 IMU 的 micro_eval，用户要求“只用轮速计”后才换成 /odom 单方向方案。")
h("8.2 问题本身的客观难度", 2)
bullet("强非线性 + 隐藏状态：stick-slip（静摩擦≫动摩擦）+ 被动万向轮角度不可观，行为高度依赖历史，单次实验难以收敛。")
bullet("观测/驱动双重受限：无单轮编码器、无力矩/电流接口、仅 chassis 级 /odom，可辨识量天然少。")
bullet("关键信息属用户领域知识、无法先验获知：增量式 PI、PI 默认值 300/300→1200、已有抗饱和、脚轮推不动、2×2 场地——这些每暴露一条都重塑方案。")
bullet("约束是逐步给出的：接口、传感、场地、控制器形式分多轮才明确，每条新约束都推翻或收窄了上一轮结论。")
h("8.3 下次更快到位的做法", 2)
bullet("开局先做“能力边界问询”：控制接口（速度/力矩/位置）、反馈（单轮?IMU?）、控制器形式与现值、场地、能否手动复位——再给方案。")
bullet("先用最小实验测 plant 特性，再下结论；不靠先验教条。")
bullet("永远先看原始时序波形，再信汇总指标；非线性 plant 上指标尤其会骗人。")
bullet("每个实验设计先问“被测量会不会被副作用污染”（如脚轮翻向、未到稳态、工具依赖项）。")

out = os.path.join(BASE, "PID调参报告_20260608.docx")
doc.save(out)
print("生成:", out)
